"""Convert uploaded files to Markdown sections (one per page, slide, heading or block of rows).

Sections are what the user sees and edits in the document view, and what gets chunked and embedded.
Pages and slides without readable text go through OCR (see ocr.py), depending on OCR_MODE.
"""

import csv
import datetime as dt
import io
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from pathlib import Path

from . import nvidia, ocr
from .config import get_settings

FILE_TYPES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".csv": "csv",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
    ".bmp": "image",
    ".tif": "image",
    ".tiff": "image",
}

# Old binary Office formats need a converter we don't ship; tell the user what to do instead.
LEGACY_FORMATS = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}

# Sections longer than this are split so they stay comfortable to read and edit.
MAX_SECTION_CHARS = 6000

Progress = Callable[[str], None]


class ExtractionError(Exception):
    """A problem with the file that the user can fix (wrong format, no text, password...)."""


@dataclass
class Section:
    markdown: str
    location: str | None = None
    page: int | None = None
    origin: str = "text"  # text | ocr
    notice: str | None = None


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in FILE_TYPES:
        return FILE_TYPES[ext]
    if ext in LEGACY_FORMATS:
        raise ExtractionError(
            f"{ext} files are an older Office format. Open the file and save it as {LEGACY_FORMATS[ext]}, then upload again."
        )
    supported = ", ".join(sorted(FILE_TYPES))
    raise ExtractionError(f"Unsupported file type '{ext or filename}'. Supported types: {supported}.")


def convert(path: Path, file_type: str, progress: Progress) -> list[Section]:
    max_chars = get_settings().chunk_size
    converter = {
        "pdf": lambda: _convert_pdf(path, progress),
        "docx": lambda: _split_long(_convert_docx(path)),
        "pptx": lambda: _convert_pptx(path, progress),
        "xlsx": lambda: _convert_xlsx(path, max_chars),
        "csv": lambda: _table_sections(csv.reader(_read_text(path).splitlines()), path.name, max_chars),
        "txt": lambda: _split_long([Section(_read_text(path))], part_label="Part"),
        "md": lambda: _split_long(_markdown_sections(_read_text(path))),
        "image": lambda: _convert_image(path, progress),
    }[file_type]
    sections = [
        Section(_clean(s.markdown), s.location, s.page, s.origin, s.notice) for s in converter()
    ]
    # PDF pages are kept even when empty so page numbers line up and the page can be OCR'd later.
    return [s for s in sections if s.markdown or s.page is not None]


def ocr_failure_notice(exc: Exception) -> str:
    return f"OCR didn't work on this page ({nvidia.describe_api_error(exc)})."


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _visible_chars(text: str) -> int:
    return len(re.sub(r"\s", "", text))


# --------------------------------------------------------------------------- PDF


def _convert_pdf(path: Path, progress: Progress) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    settings = get_settings()
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("This PDF is password-protected. Remove the password and upload it again.")
        total = len(reader.pages)
        texts = []
        for number, page in enumerate(reader.pages, start=1):
            try:
                texts.append(page.extract_text() or "")
            except Exception:  # a single broken page shouldn't sink the document
                texts.append("")
            if number % 20 == 0:
                progress(f"Reading page {number} of {total}")
    except PdfReadError as exc:
        raise ExtractionError(f"Couldn't read this PDF: {exc}") from exc

    if settings.ocr_mode == "off":
        to_ocr: list[int] = []
    else:
        to_ocr = [
            number
            for number, text in enumerate(texts, start=1)
            if settings.ocr_mode == "always" or _visible_chars(text) < settings.ocr_min_text_chars
        ]
    results = (
        ocr.ocr_pdf_pages(path, to_ocr, lambda done, n: progress(f"Running OCR: page {done} of {n}"))
        if to_ocr
        else {}
    )

    sections = []
    for number, text in enumerate(texts, start=1):
        result = results.get(number)
        location = f"Page {number}"
        if isinstance(result, str) and result.strip():
            sections.append(Section(result, location, page=number, origin="ocr"))
            continue
        notice = None
        if isinstance(result, Exception):
            notice = ocr_failure_notice(result) + (" Showing the text stored in the PDF instead." if text.strip() else "")
        elif result == "" and not text.strip():
            notice = "OCR found no text on this page."
        sections.append(Section(_pdf_text_to_markdown(text), location, page=number, notice=notice))

    if not any(s.markdown.strip() for s in sections):
        if settings.ocr_mode == "off":
            raise ExtractionError("No text found in this PDF. It's probably scanned images; set OCR_MODE=auto to read it with OCR.")
        failures = [r for r in results.values() if isinstance(r, Exception)]
        if failures:
            raise ExtractionError(f"No text found in this PDF. {ocr_failure_notice(failures[0])}")
        raise ExtractionError("No text found in this PDF, even with OCR.")
    return sections


def _pdf_text_to_markdown(text: str) -> str:
    """Re-join the hard-wrapped lines of PDF text into paragraphs.

    A line ends a paragraph when it's noticeably shorter than the page's usual line length,
    which is how headings, list items and paragraph ends look in extracted text.
    """
    lines = [line.strip() for line in text.splitlines()]
    lengths = sorted(len(line) for line in lines if line)
    if not lengths:
        return ""
    typical = lengths[int(len(lengths) * 0.8)]

    paragraphs: list[str] = []
    current = ""
    for line in lines:
        if not line:
            if current:
                paragraphs.append(current)
                current = ""
            continue
        if current.endswith("-") and line[:1].islower():
            current = current[:-1] + line  # re-join a hyphenated word
        else:
            current = f"{current} {line}" if current else line
        if len(line) < typical * 0.7:
            paragraphs.append(current)
            current = ""
    if current:
        paragraphs.append(current)
    return "\n\n".join(paragraphs)


# --------------------------------------------------------------------------- Images


def _convert_image(path: Path, progress: Progress) -> list[Section]:
    from PIL import Image, UnidentifiedImageError

    if get_settings().ocr_mode == "off":
        raise ExtractionError("Images can only be read with OCR, which is off. Set OCR_MODE=auto in backend/.env.")
    try:
        image = Image.open(path)
        image.load()
    except (UnidentifiedImageError, OSError) as exc:
        raise ExtractionError(f"Couldn't open this image: {exc}") from exc
    progress("Running OCR on the image")
    markdown = ocr.image_to_markdown(image)
    if not markdown.strip():
        raise ExtractionError("OCR found no text in this image.")
    return [Section(markdown, None, origin="ocr")]


# --------------------------------------------------------------------------- Word


def _convert_docx(path: Path) -> list[Section]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    sections: list[Section] = []
    heading: str | None = None
    blocks: list[str] = []

    def flush() -> None:
        if blocks:
            sections.append(Section(_join_blocks(blocks), heading))
            blocks.clear()

    for block in document.iter_inner_content():
        if isinstance(block, Table):
            blocks.append(_markdown_table([[cell.text for cell in row.cells] for row in block.rows]))
            continue
        if not isinstance(block, Paragraph) or not block.text.strip():
            continue
        style = (block.style.name if block.style is not None else "") or ""
        level = _heading_level(style)
        if level:
            flush()
            heading = block.text.strip()[:80]
            blocks.append(f"{'#' * level} {block.text.strip()}")
        else:
            blocks.append(_list_prefix(block, style) + _runs_markdown(block))
    flush()
    return sections


def _heading_level(style: str) -> int:
    if style == "Title":
        return 1
    match = re.match(r"Heading (\d)", style)
    return min(int(match.group(1)), 6) if match else 0


def _list_prefix(paragraph, style: str) -> str:
    p_pr = paragraph._p.pPr
    num_pr = p_pr.numPr if p_pr is not None else None
    if not (style.startswith("List") or num_pr is not None):
        return ""
    level = 0
    if num_pr is not None and num_pr.ilvl is not None:
        level = int(num_pr.ilvl.val)
    elif match := re.search(r"(\d)$", style):
        level = int(match.group(1)) - 1
    marker = "1." if "Number" in style else "-"
    return "  " * level + marker + " "


def _runs_markdown(paragraph) -> str:
    """Paragraph text with bold/italic and links kept as Markdown."""
    from docx.text.hyperlink import Hyperlink

    pieces: list[tuple[str, bool, bool]] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            text = item.text
            pieces.append((f"[{text}]({item.url})" if item.url and text.strip() else text, False, False))
        else:
            pieces.append((item.text, bool(item.bold), bool(item.italic)))

    merged: list[list] = []
    for text, bold, italic in pieces:
        if merged and merged[-1][1:] == [bold, italic]:
            merged[-1][0] += text
        else:
            merged.append([text, bold, italic])

    out = []
    for text, bold, italic in merged:
        core = text.strip()
        if not core or not (bold or italic):
            out.append(text)
            continue
        mark = "***" if bold and italic else "**" if bold else "*"
        lead = text[: len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()) :]
        out.append(f"{lead}{mark}{core}{mark}{trail}")
    return "".join(out).strip()


def _join_blocks(blocks: list[str]) -> str:
    """Blank lines between blocks, except between consecutive list items."""
    out = ""
    prev_item = False
    for block in blocks:
        is_item = bool(re.match(r"\s*([-*]|\d+\.)\s", block))
        if out:
            out += "\n" if is_item and prev_item else "\n\n"
        out += block
        prev_item = is_item
    return out


# --------------------------------------------------------------------------- PowerPoint


def _convert_pptx(path: Path, progress: Progress) -> list[Section]:
    from pptx import Presentation

    settings = get_settings()
    presentation = Presentation(str(path))
    sections = []
    total = len(presentation.slides)
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        lines = [f"## {title}"] if title else []
        pictures: list = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            lines.extend(_shape_markdown(shape, pictures))

        origin, notice = "text", None
        body = "\n".join(lines)
        # Slides that are mostly a screenshot or scanned image: read the picture text with OCR.
        if settings.ocr_mode != "off" and pictures and _visible_chars(body) < settings.ocr_min_text_chars:
            progress(f"Running OCR on slide {number} of {total}")
            for blob in sorted(pictures, key=len, reverse=True)[:3]:
                try:
                    from PIL import Image

                    text = ocr.image_to_markdown(Image.open(io.BytesIO(blob)))
                except ocr.FATAL_ERRORS:
                    raise
                except Exception as exc:  # unreadable image format or OCR failure
                    notice = ocr_failure_notice(exc)
                    continue
                if text.strip():
                    lines.append(text)
                    origin = "ocr"

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("\n> **Speaker notes:** " + notes.replace("\n", "\n> "))

        location = f"Slide {number}" + (f": {title[:60]}" if title else "")
        sections.append(Section(_join_blocks(lines) if lines else "", location, origin=origin, notice=notice))
    return sections


def _shape_markdown(shape, pictures: list) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for child in shape.shapes for line in _shape_markdown(child, pictures)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [_markdown_table([[cell.text for cell in row.cells] for row in shape.table.rows])]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [
            "  " * p.level + "- " + p.text.strip()
            for p in shape.text_frame.paragraphs
            if p.text.strip()
        ]
    try:
        pictures.append(shape.image.blob)
    except (AttributeError, ValueError):
        pass
    return []


# --------------------------------------------------------------------------- Spreadsheets


def _convert_xlsx(path: Path, max_chars: int) -> list[Section]:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a variety of zip/xml errors
        raise ExtractionError(f"Couldn't open this spreadsheet: {exc}") from exc
    try:
        sections = []
        for sheet in workbook.worksheets:
            sections.extend(_table_sections(sheet.iter_rows(values_only=True), f"Sheet {sheet.title}", max_chars))
        return sections
    finally:
        workbook.close()


def _table_sections(rows: Iterable[Iterable], label: str, max_chars: int) -> list[Section]:
    """Markdown tables of consecutive rows, each small enough to be one search chunk.

    Every section repeats the header row, so a retrieved chunk still says what each column is.
    """
    header: list[str] | None = None
    sections: list[Section] = []
    body: list[list[str]] = []
    first_row = last_row = 0

    def flush() -> None:
        if body and header is not None:
            rows_label = f"row {first_row}" if first_row == last_row else f"rows {first_row}-{last_row}"
            sections.append(Section(_markdown_table([header, *body]), f"{label}, {rows_label}"))
            body.clear()

    size = header_size = 0
    for row_number, row in enumerate(rows, start=1):
        values = [_cell_str(v) for v in row]
        while values and not values[-1]:
            values.pop()
        if not values:
            continue
        if header is None:
            header = [v or f"Column {i + 1}" for i, v in enumerate(values)]
            # The header and separator lines are repeated in every section.
            header_size = 2 * sum(len(h) + 3 for h in header)
            continue
        row_size = sum(len(v) + 3 for v in values) + 2
        if body and header_size + size + row_size > max_chars:
            flush()
            size = 0
        if not body:
            first_row = row_number
        body.append(values)
        size += row_size
        last_row = row_number
    flush()

    if header is not None and not sections:
        sections.append(Section(" | ".join(h for h in header if h), label))
    return sections


def _markdown_table(rows: list[list[str]]) -> str:
    rows = [[" ".join(str(c).split()).replace("|", "\\|") for c in row] for row in rows if any(str(c).strip() for c in row)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()


# --------------------------------------------------------------------------- Text and Markdown


def _markdown_sections(text: str) -> list[Section]:
    sections: list[Section] = []
    heading: str | None = None
    lines: list[str] = []
    in_code = False
    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
        match = None if in_code else re.match(r"#{1,3}\s+(.+)", line)
        if match and any(l.strip() for l in lines):
            sections.append(Section("\n".join(lines), heading))
            lines = []
        if match:
            heading = match.group(1).strip()[:80]
        lines.append(line)
    if any(l.strip() for l in lines):
        sections.append(Section("\n".join(lines), heading))
    return sections


def _split_long(sections: list[Section], part_label: str | None = None) -> list[Section]:
    """Split very long sections at paragraph breaks so each stays easy to read and edit."""
    out: list[Section] = []
    for section in sections:
        paragraphs = section.markdown.split("\n\n")
        parts: list[str] = []
        current = ""
        for paragraph in paragraphs:
            if current and len(current) + len(paragraph) > MAX_SECTION_CHARS:
                parts.append(current)
                current = ""
            current = f"{current}\n\n{paragraph}" if current else paragraph
        if current:
            parts.append(current)

        for i, part in enumerate(parts, start=1):
            if len(parts) == 1:
                location = section.location
            elif part_label:
                location = f"{part_label} {i}"
            else:
                location = f"{section.location or 'Part'} ({i}/{len(parts)})"
            out.append(Section(part, location, section.page, section.origin, section.notice))
    return out


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")
